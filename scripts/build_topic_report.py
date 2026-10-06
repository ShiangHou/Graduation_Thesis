"""Build the advisor report from its version-controlled Markdown source."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'reports' / '硕士论文选题报告.md'
OUTPUT = ROOT / 'reports' / '硕士论文选题报告.docx'

def set_fonts(style, western='Calibri', east='Songti SC'):
    style.font.name = western
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:' + attr), None)
    fonts.set(qn('w:eastAsia'), east)
    fonts.set(qn('w:ascii'), western)
    fonts.set(qn('w:hAnsi'), western)
    style.font.color.rgb = RGBColor(0, 0, 0)

def hyperlink(p, value):
    rid = p.part.relate_to(value, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '174A73')
    props.append(color)
    run.append(props)
    t = OxmlElement('w:t')
    t.text = value
    run.append(t)
    link.append(run)
    p._p.append(link)

def inline(p, value):
    for token in re.split(r'(https?://[^\s]+|\*\*[^*]+\*\*)', value):
        if token.startswith('http'):
            hyperlink(p, token)
        elif token.startswith('**') and token.endswith('**'):
            p.add_run(token[2:-2]).bold = True
        else:
            p.add_run(token)

def table(doc, lines):
    rows = [[v.strip() for v in line.strip().strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r':?-+:?', c) for c in r)]
    count = len(rows[0])
    t = doc.add_table(rows=0, cols=count)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    width = 17.0
    proportions = {3:[.23,.37,.40],4:[.18,.30,.24,.28],5:[.19,.26,.19,.22,.14],6:[.24,.18,.23,.13,.13,.09]}.get(count,[1/count]*count)
    for col, part in zip(t.columns, proportions): col.width = Cm(width*part)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        e = OxmlElement('w:'+name)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]: e.set(qn('w:'+k),v)
        borders.append(e)
    pr.append(borders)
    for i, items in enumerate(rows):
        row = t.add_row()
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trpr.append(cant)
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        for j, value in enumerate(items):
            cell = row.cells[j]
            cell.width = Cm(width*proportions[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                item=OxmlElement('w:'+side); item.set(qn('w:w'),'90'); item.set(qn('w:type'),'dxa'); margins.append(item)
            tcpr.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'),'E8EDF2' if i==0 else ('F7F8FA' if i%2==0 else 'FFFFFF'))
            tcpr.append(shade)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.space_before=Pt(2)
            p.paragraph_format.line_spacing=1.10
            inline(p,value)
            for r in p.runs:
                r.font.size=Pt(9)
                if i==0: r.bold=True
    spacer=doc.add_paragraph()
    spacer.paragraph_format.space_after=Pt(2)
    spacer.paragraph_format.space_before=Pt(0)
    spacer.paragraph_format.line_spacing=0.5
    spacer.add_run().font.size=Pt(2)

def main():
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.0); sec.bottom_margin=Cm(1.9)
    sec.left_margin=Cm(2.0); sec.right_margin=Cm(2.0)
    sec.footer_distance=Cm(.8)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet']:
        set_fonts(doc.styles[name], east='Songti SC')
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    normal=doc.styles['Normal']
    normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.20
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.widow_control=True
    doc.styles['Title'].font.size=Pt(21)
    doc.styles['Title'].font.bold=True
    doc.styles['Title'].paragraph_format.space_after=Pt(16)
    for name,size in [('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
        s=doc.styles[name]; s.font.size=Pt(size); s.font.bold=True
        s.paragraph_format.space_before=Pt(11);s.paragraph_format.space_after=Pt(6)
        s.paragraph_format.keep_with_next=True
    doc.core_properties.title='计算机技术与金融信息检索硕士论文选题报告'
    doc.core_properties.subject='三个实习衍生选题的研究问题 文献 实验与可行性'
    doc.core_properties.author=''
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    lines=SOURCE.read_text().splitlines()
    i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1;continue
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i]);i+=1
            table(doc,block);continue
        if line.startswith('# '):
            doc.add_paragraph(line[2:],style='Title')
        elif line.startswith('## '):
            p=doc.add_paragraph(line[3:],style='Heading 1')
            if line.startswith(('## 三 ','## 四 ','## 五 ')):
                p.paragraph_format.page_break_before=True
        elif line.startswith('### '):
            doc.add_paragraph(line[4:],style='Heading 2')
        elif line.startswith('- '):
            inline(doc.add_paragraph(style='List Bullet'),line[2:])
        else:
            p=doc.add_paragraph();inline(p,line)
            if re.match(r'^\[\d+\]',line):
                p.paragraph_format.space_after=Pt(6)
                for r in p.runs:r.font.size=Pt(9)
        i+=1
    doc.save(OUTPUT)
    print(OUTPUT)
    print('Source characters:',len(SOURCE.read_text()))
    print('Paragraphs:',len(doc.paragraphs),'Tables:',len(doc.tables))

if __name__=='__main__':main()
